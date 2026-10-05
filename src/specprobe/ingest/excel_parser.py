from openpyxl import load_workbook

from .base import TextBlock


class ExcelDidTableParser:
    def parse(self, path: str) -> list[TextBlock]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks: list[TextBlock] = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            previous: list[object] = []
            for row in sheet.iter_rows(values_only=True):
                values = list(row)
                if any(value is not None for value in values):
                    previous = [value if value is not None else previous[pos] if pos < len(previous) else "" for pos, value in enumerate(values)]
                    blocks.append(TextBlock(" | ".join(str(value) for value in previous), index))
        return blocks
