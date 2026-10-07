import logging
import re
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber

from .layouts import header as normalize_header
from .layouts import is_header
from .models import TableRow, TextBlock

LOGGER = logging.getLogger(__name__)


def parse_excel(path: str | Path, document_id: str) -> tuple[list[TextBlock], list[TableRow]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    blocks: list[TextBlock] = []
    rows: list[TableRow] = []
    for sheet in workbook.worksheets:
        merged_values: dict[tuple[int, int], Any] = {}
        for merged in sheet.merged_cells.ranges:
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    merged_values[(row, column)] = value
        headers: list[str] | None = None
        first_row = 1
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=first_row, values_only=True), first_row
        ):
            values = [
                merged_values.get((row_number, column), value)
                for column, value in enumerate(cells, 1)
            ]
            if not any(value is not None for value in values):
                continue
            prose = " ".join(str(value) for value in values if value is not None)
            if prose:
                blocks.append(TextBlock(document_id, prose, sheet=sheet.title, row=row_number))
            if sheet.title == "Appendix":
                continue
            if headers is None:
                if not is_header(values):
                    continue
                headers = [str(value) for value in values]
                continue
            if [normalize_header(value) for value in values] == [
                normalize_header(value) for value in headers
            ]:
                continue
            record = {
                header: values[index]
                for index, header in enumerate(headers)
                if index < len(values) and header
            }
            rows.append(TableRow(document_id, record, sheet=sheet.title, row=row_number))
    workbook.close()
    if not rows and any(
        re.search(r"\b(?:DID|routine|session|security)\b", block.text, re.I) for block in blocks
    ):
        raise ValueError(
            "Unsupported diagnostic spreadsheet layout; supply a reviewed EcuSpec JSON"
        )
    return blocks, rows


def parse_pdf(path: str | Path, document_id: str) -> tuple[list[TextBlock], list[TableRow]]:
    blocks: list[TextBlock] = []
    rows: list[TableRow] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(
                    f"Page {page_number} has no extractable text; OCR and engineer review required"
                )
            blocks.append(TextBlock(document_id, text, page=page_number))
            tables = page.extract_tables()
            for table_index, table in enumerate(tables):
                if not table or not table[0]:
                    continue
                header_row = table[0]
                data_start = 1
                if len([value for value in header_row if value]) == 1 and len(table) > 1:
                    header_row = table[1]
                    data_start = 2
                headers = [str(value or "").strip() for value in header_row]
                if not all(headers):
                    LOGGER.warning(
                        "ambiguous table: document=%s page=%s table=%s",
                        document_id,
                        page_number,
                        table_index,
                    )
                    raise ValueError(f"Ambiguous PDF table on page {page_number}")
                for row_number, values in enumerate(table[data_start:], data_start + 1):
                    if len(values) != len(headers):
                        LOGGER.warning(
                            "ambiguous table row: document=%s page=%s row=%s",
                            document_id,
                            page_number,
                            row_number,
                        )
                        raise ValueError(
                            f"Ambiguous PDF table row on page {page_number}: {row_number}"
                        )
                    rows.append(
                        TableRow(
                            document_id,
                            dict(zip(headers, values, strict=True)),
                            page=page_number,
                            row=row_number,
                            table=table_index + 1,
                        )
                    )
    return blocks, rows
